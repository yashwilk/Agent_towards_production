"""Generates a realistic test docx (a vendor security review report, complete with
a real Word table) for exercising the conversion client without needing an input file."""

from docx import Document

import config


def write_sample_docx(path: str = config.SAMPLE_DOCX_PATH) -> None:
    doc = Document()
    doc.add_heading("Vendor Security Review, Q3", level=1)
    doc.add_paragraph(
        "Four vendors were reviewed this quarter. One review remains open pending "
        "an updated penetration-test report."
    )

    table = doc.add_table(rows=5, cols=3)
    table.style = "Table Grid"
    rows = [
        ("Vendor", "Risk score", "Status"),
        ("Acme Cloud", "Low", "Approved"),
        ("DataBridge", "Medium", "Approved with conditions"),
        ("Sky CRM", "High", "Rejected"),
        ("OptiLog", "Medium", "Open"),
    ]
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.rows[r].cells[c].text = text

    doc.save(path)
    print(f"Wrote {path}")
